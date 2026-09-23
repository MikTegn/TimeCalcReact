#!/usr/bin/env python3
"""
Konverterar TimeCalc-arbetsboken (.xlsm) till JSON som React-appen läser.

    python3 scripts/import_xlsm.py ~/Documents/Arbete/TimeCalc-25.xlsm

Skriver:
  data/timecalc-<år>.json   indata (aktiviteter, veckor, dagar, tidslogg, rapporterade timmar)
  data/expected-<år>.json   Excels egna (cachade) beräknade värden, används av tests/golden.test.ts
  public/seed/timecalc.json kopia av indata som appen laddar första gången (hoppa över med --no-seed)

Kräver: python3 och openpyxl (pip install openpyxl). Makron (VBA) körs aldrig.
"""
import argparse, collections, datetime as dt, json, re, shutil, sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as CI

# Arbetsbladets layout (se README): tio dagblock à tre kolumner. Bara mån–fre (block 0–4) används.
BLOCK_FIRST_COLS = [CI(c) for c in ("G", "J", "M", "P", "S", "V", "Y", "AB", "AE", "AH")]
DAYS_IN_MODEL = 5
ACT_ROWS = range(4, 22)      # D4:D21 – aktivitetslistan (18 st)
LOG_START_ROW = 24           # H24 – dagens starttid
LOG_ROWS = range(25, 46)     # G25:H45 – aktivitet + sluttid (21 rader)
MAX_ROW, MAX_COL = 48, 39    # läs aldrig längre: V36 har felaktig dimension (1 048 576 rader)

warnings = []


def warn(msg):
    warnings.append(msg)
    print("  ! " + msg)


def to_minutes(v):
    """Excel-tid (time/timedelta/datetime/serietal) -> minuter, eller None för tom cell."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return round((v - dt.datetime(1899, 12, 30)).total_seconds() / 60, 6)
    if isinstance(v, dt.time):
        return round(v.hour * 60 + v.minute + v.second / 60, 6)
    if isinstance(v, dt.timedelta):
        return round(v.total_seconds() / 60, 6)
    if isinstance(v, (int, float)):
        return round(v * 24 * 60, 6)
    return None


def to_number(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def eval_simple_formula(f):
    """Utvärderar rena aritmetiska formler som '=38.5/5'. Allt annat -> None."""
    body = f.lstrip("=").strip()
    if re.fullmatch(r"[\d.+\-*/() ]+", body):
        try:
            return float(eval(body, {"__builtins__": {}}, {}))  # noqa: S307 – tecken verifierade ovan
        except Exception:
            return None
    return None


def describe_value(v):
    if isinstance(v, dt.timedelta):
        minutes = int(v.total_seconds() // 60)
        return f"{minutes // 60}:{minutes % 60:02d}"
    if isinstance(v, dt.time):
        return f"{v.hour}:{v.minute:02d}"
    return repr(v)


def note_of(*cells):
    parts = [c.comment.text.strip() for c in cells if c is not None and c.comment and c.comment.text.strip()]
    return "\n".join(parts) or None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsm")
    ap.add_argument("--out", default="data")
    ap.add_argument("--no-seed", action="store_true")
    args = ap.parse_args()

    src = Path(args.xlsm).expanduser()
    if not src.exists():
        sys.exit(f"Hittar inte {src}")
    print(f"Läser {src} ...")
    wf = load_workbook(src)                    # formler
    wv = load_workbook(src, data_only=True)    # Excels cachade värden

    ov_f, ov_v = wf["Overview"], wv["Overview"]

    # ---- aktiviteter (Overview!G2:X2) ----------------------------------------------------
    activities = []
    for i, r in enumerate(ACT_ROWS):
        col = 7 + i  # G..X
        name = ov_f.cell(2, col).value
        if not isinstance(name, str) or not name.strip():
            sys.exit(f"Overview {ov_f.cell(2, col).coordinate} saknar aktivitetsnamn")
        activities.append({"name": name, "code": note_of(ov_f.cell(2, col)), "_col": col, "_row": r})
    names = [a["name"] for a in activities]
    if len({n.lower() for n in names}) != len(names):
        warn("Aktivitetsnamn är inte unika (SUMIF ignorerar versaler) – summeringar kan dubbelräknas")

    # ---- veckoblad i Overview-ordning -----------------------------------------------------
    weeks_meta = []
    r = 4
    while ov_f.cell(r, 5).value:                       # E4.. = veckoetikett
        label = str(ov_f.cell(r, 5).value)
        m = re.fullmatch(r"='?([^'!]+)'?!B23", str(ov_f.cell(r, 1).value or ""))
        if not m:
            sys.exit(f"Overview A{r}: väntade ='Vxx'!B23, fick {ov_f.cell(r, 1).value!r}")
        weeks_meta.append({"label": label, "sheet": m.group(1), "ov_row": r})
        r += 1
    print(f"  {len(weeks_meta)} veckor, {len(activities)} aktiviteter")

    # ---- konfiguration --------------------------------------------------------------------
    first = wf[weeks_meta[0]["sheet"]]
    m = re.fullmatch(r"=([\d.]+)\*COUNTIF\(.*\)/(\d+)", str(first["B23"].value))
    weekly_hours, per_week = (float(m.group(1)), int(m.group(2))) if m else (38.5, 5)
    m = re.fullmatch(r"=([\d.]+)\+C24", str(first["C25"].value))
    opening = float(m.group(1)) if m else 0.0
    year = wf[weeks_meta[1]["sheet"]]["G2"].value.year  # V02 ligger säkert inom året

    # ---- veckor --------------------------------------------------------------------------
    weeks, flag_counter = [], collections.defaultdict(collections.Counter)
    for wm in weeks_meta:
        ws = wf[wm["sheet"]]
        days = []
        for bi, c0 in enumerate(BLOCK_FIRST_COLS):
            date = ws.cell(2, c0).value
            flag_cell, hdr_cell = ws.cell(3, c0 + 2), ws.cell(3, c0 + 1)
            log = [(ws.cell(r, c0), ws.cell(r, c0 + 1)) for r in LOG_ROWS]
            used = (
                bool(flag_cell.value)
                or any(ws.cell(r, c0 + 2).value not in (None, "") for r in range(4, 21))
                or any(a.value not in (None, "") or t.value not in (None, "") for a, t in log)
                or ws.cell(LOG_START_ROW, c0 + 1).value not in (None, "")
            )
            if bi >= DAYS_IN_MODEL:
                if used:
                    warn(f"{wm['sheet']}: dagblock {bi + 1} ({date:%Y-%m-%d}) innehåller data men ingår inte i modellen (bara mån–fre)")
                continue

            rows = []
            for a, t in log:
                act = a.value if isinstance(a.value, str) else ""
                end = to_minutes(t.value)
                if t.value not in (None, "") and end is None:
                    warn(f"{wm['sheet']}!{t.coordinate}: kan inte tolka tid {t.value!r}")
                rows.append({"activity": act, "end": end, "note": note_of(a, t)})
            while rows and not rows[-1]["activity"] and rows[-1]["end"] is None and not rows[-1]["note"]:
                rows.pop()
            for row in rows:
                if not row["note"]:
                    del row["note"]

            reported = {}
            for i, r in enumerate(range(4, 21)):       # I4:I20 – manuellt rapporterade timmar
                cell = ws.cell(r, c0 + 2)
                v = cell.value
                if v in (None, ""):
                    continue
                num = to_number(v)
                if num is None and isinstance(v, str) and v.startswith("="):
                    num = eval_simple_formula(v)
                if num is None:
                    warn(f"{wm['sheet']}!{cell.coordinate}: ogiltigt värde {v!r} i rapporterade timmar ({names[i]}) – ignoreras")
                    continue
                reported[names[i]] = num

            days.append({
                "date": date.date().isoformat(),
                "workday": str(flag_cell.value or "").strip().lower() == "ja",
                "forceRegistered": isinstance(hdr_cell.value, str) and not hdr_cell.value.startswith("=") and hdr_cell.value.strip().lower() == "tid",
                "start": to_minutes(ws.cell(LOG_START_ROW, c0 + 1).value),
                "rows": rows,
                "reported": reported,
            })

        flags = {}
        for i, r in enumerate(ACT_ROWS):
            f = {"report": bool(ws.cell(r, 1).value), "bank": bool(ws.cell(r, 2).value), "normal": bool(ws.cell(r, 3).value)}
            flags[names[i]] = f
            for k, v in f.items():
                flag_counter[names[i]][(k, v)] += 1
        weeks.append({"id": wm["sheet"], "label": wm["label"], "days": days, "flags": flags})

    # ---- avvikelser i arbetsboken: formler som skrivits över med konstanter eller raderats -----
    # En cell som har exakt samma formel i nästan alla veckoblad men är en konstant (eller tom) i något
    # blad är sannolikt ett misstag. Excel räknar då vidare med det felaktiga värdet.
    grids = {}
    for wm in weeks_meta:
        ws = wf[wm["sheet"]]
        grids[wm["sheet"]] = {
            c.coordinate: c.value
            for row in ws.iter_rows(min_row=1, max_row=MAX_ROW, max_col=MAX_COL)
            for c in row
            if c.value is not None
        }
    is_formula = lambda v: isinstance(v, str) and v.startswith("=")
    coords = {k for g in grids.values() for k, v in g.items() if is_formula(v)}
    anomalies = []
    for coord in sorted(coords, key=lambda k: (int(re.sub(r"\D", "", k)), k)):
        texts = collections.Counter(g.get(coord) for g in grids.values() if is_formula(g.get(coord)))
        if not texts:
            continue
        common, n = texts.most_common(1)[0]
        if n < 0.85 * len(grids):
            continue  # formeln är avsiktligt olika mellan blad (t.ex. B23, C25)
        for sheet, g in grids.items():
            v = g.get(coord)
            if not is_formula(v):  # konstant eller tom – formler som skiljer sig från varandra räknas som varianter
                anomalies.append({"sheet": sheet, "cell": coord, "expectedFormula": common, "found": None if v is None else describe_value(v)})
    for a in anomalies:
        what = "saknas" if a["found"] is None else f"har konstanten {a['found']}"
        warn(f"{a['sheet']}!{a['cell']}: {what} i stället för formeln {a['expectedFormula']}")

    # standardflaggor = vanligaste värdet över alla veckor
    for a in activities:
        a["defaults"] = {k: flag_counter[a["name"]][(k, True)] >= flag_counter[a["name"]][(k, False)] for k in ("report", "bank", "normal")}
        # Rapporterade timmar för Lunch härleds ur loggad tid (I21 = HOUR(H21)+MINUTE(H21)/60) i stället för att matas in
        formula = str(first.cell(a["_row"], 9).value or "")
        a["derivedReport"] = bool(re.search(r"HOUR\(H\d+\)\+MINUTE\(H\d+\)/60", formula))
        del a["_row"], a["_col"]

    dataset = {
        "version": 1,
        "source": {"file": src.name, "importedAt": dt.datetime.now().isoformat(timespec="seconds")},
        "config": {"year": year, "weeklyHours": weekly_hours, "workdaysPerWeek": per_week, "openingBank": opening},
        "activities": activities,
        "weeks": weeks,
    }

    # ---- Excels cachade värden (facit för tests) ------------------------------------------
    def hrs(cell):
        return to_number(cell.value)

    exp_weeks = {}
    for wm in weeks_meta:
        v = wv[wm["sheet"]]
        acts = {}
        for i, r in enumerate(ACT_ROWS):
            acts[names[i]] = {"registeredMin": to_minutes(v.cell(r, 5).value), "reportedHours": to_number(v.cell(r, 6).value)}
        days = []
        for c0 in BLOCK_FIRST_COLS[:DAYS_IN_MODEL]:
            days.append({
                "summedMin": to_minutes(v.cell(3, c0).value),
                "rptRegisteredMin": to_minutes(v.cell(22, c0 + 1).value),
                "rptReportedHours": to_number(v.cell(22, c0 + 2).value),
            })
        exp_weeks[wm["sheet"]] = {
            "bankHours": hrs(v["B22"]), "expectedHours": hrs(v["B23"]), "bankDiff": hrs(v["B24"]),
            "normalHours": hrs(v["C22"]), "normalDiff": hrs(v["C24"]), "normalBalance": hrs(v["C25"]),
            "rptHours": hrs(v["F22"]), "rptRegisteredMin": to_minutes(v["E22"].value),
            "registeredMin": to_minutes(v["E32"].value), "summedMin": to_minutes(v["E33"].value),
            "unmappedMin": to_minutes(v["E34"].value), "reportedHours": hrs(v["E35"]),
            "reportedAsTimeMin": to_minutes(v["E36"].value), "regVsReportedMin": to_minutes(v["E37"].value),
            "activities": acts, "days": days,
        }

    # ---- Excels varningar (villkorsformat) utvärderade på de cachade värdena ---------------
    # Reglerna är avlästa ur arbetsbokens villkorsformatering (se README). Fel i en regelformel
    # (t.ex. DAY("") när en cell är tom) ger ingen markering, precis som i Excel.
    def emulate_warnings(sheet):
        v, f = wv[sheet], wf[sheet]
        hrs_of = lambda cell: None if to_minutes(cell.value) is None else to_minutes(cell.value) / 60
        out = []
        for bi, c0 in enumerate(BLOCK_FIRST_COLS[:DAYS_IN_MODEL]):
            for r in range(4, 21):                                   # I4:I20 mot H4:H20
                h = hrs_of(v.cell(r, c0 + 1))
                if h is None:
                    continue
                raw = f.cell(r, c0 + 2).value
                if raw is None:
                    i = 0.0                                          # tom cell räknas som 0
                elif isinstance(raw, str) and raw.startswith("="):
                    i = to_number(v.cell(r, c0 + 2).value)
                else:
                    i = to_number(raw)                               # text ger #VALUE! -> ingen markering
                if i is not None and abs(i - h) > 0.25 + 1e-9:
                    out.append(f"dayActivity:{bi}:{names[r - 4]}")
            i22, h22 = to_number(v.cell(22, c0 + 2).value), hrs_of(v.cell(22, c0 + 1))
            if i22 is not None and h22 is not None and abs(i22 - h22) > 0.5 + 1e-9:
                out.append(f"dayTotal:{bi}")
        for i, r in enumerate(ACT_ROWS):                             # F4:F21 mot E4:E21
            fv, ev = to_number(v.cell(r, 6).value), hrs_of(v.cell(r, 5))
            if fv is not None and ev is not None and abs(fv - ev) > 0.25 + 1e-9:
                out.append(f"activityWeek:{names[i]}")
        f22, e22 = to_number(v["F22"].value), hrs_of(v["E22"])
        if f22 is not None and e22 is not None and abs(f22 - e22) > 0.25 + 1e-9:
            out.append("rptWeek")
        if (to_number(v["C24"].value) or 0) and abs(to_number(v["C24"].value)) > 1e-9:
            out.append("normalDiff")
        if (to_minutes(v["E34"].value) or 0) and abs(to_minutes(v["E34"].value)) > 1e-9:
            out.append("unmapped")
        e37 = hrs_of(v["E37"])
        if e37 is not None and abs(e37) > 0.5 + 1e-9:
            out.append("regVsReported")
        return sorted(out)

    for wm in weeks_meta:
        exp_weeks[wm["sheet"]]["warnings"] = emulate_warnings(wm["sheet"])

    ov_rows = [{
        "sheet": wm["sheet"], "expectedHours": to_number(ov_v.cell(wm["ov_row"], 1).value),
        "bankHours": to_number(ov_v.cell(wm["ov_row"], 6).value), "bankDiff": to_number(ov_v.cell(wm["ov_row"], 2).value),
        "bankBalance": to_number(ov_v.cell(wm["ov_row"], 3).value),
    } for wm in weeks_meta]
    total_row = 4 + len(weeks_meta) + 1  # raden med summorna (58 i originalet)
    if ov_f.cell(total_row, 6).value and "SUM" in str(ov_f.cell(total_row, 6).value):
        totals = {names[i]: to_number(ov_v.cell(total_row, 7 + i).value) or 0.0 for i in range(len(names))}
        grand = to_number(ov_v.cell(total_row, 6).value)
    else:
        totals, grand = {}, None
        warn("Hittade inte summeraden i Overview – årssummor testas inte")
    expected = {"version": 1, "anomalies": anomalies, "weeks": exp_weeks, "overview": {"rows": ov_rows, "activityTotals": totals, "grandTotalHours": grand}}

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    data_path, exp_path = out / f"timecalc-{year}.json", out / f"expected-{year}.json"
    data_path.write_text(json.dumps(dataset, ensure_ascii=False, indent=1), encoding="utf-8")
    exp_path.write_text(json.dumps(expected, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Skrev {data_path} och {exp_path}")
    if not args.no_seed:
        seed = Path("public/seed"); seed.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(data_path, seed / "timecalc.json")
        print(f"Skrev {seed / 'timecalc.json'}")
    print(f"Klart med {len(warnings)} varning(ar).")


if __name__ == "__main__":
    main()
